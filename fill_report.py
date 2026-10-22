"""Fill assignment_report.docx with REAL results from running the code on this computer.

Run:  python fill_report.py
Reads  assignment_report.docx  and writes  assignment_report_filled.docx
(Re-run it on your own machine so the timings are your machine's timings.)
"""
import os
import platform
import sys
import time
from decimal import Decimal

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, "Task10"))

from milkrate import rate_per_litre, payout, validate_entry, ten_day_payment  # noqa: E402
from task10 import float_total, decimal_total  # noqa: E402

NAME = "Harshan K"
REG = "1911260183"
REPO = f"CSA08-{REG}-Assignment1"


def rejected(*args):
    try:
        validate_entry(*args)
        return False
    except ValueError:
        return True


def status(ok):
    return "Pass" if ok else "Fail"


def table_tests():
    rows = []
    for litres, fat, snf, er, ep in [("10", "3.5", "8.5", "30.00", "300.00"),
                                     ("8.5", "4.2", "8.7", "33.40", "283.90"),
                                     ("12", "3.1", "8.3", "27.80", "333.60")]:
        r, p = rate_per_litre(fat, snf), payout(litres, fat, snf)
        rows.append((f"Rate ₹{r}; pay ₹{p}", status(r == Decimal(er) and p == Decimal(ep))))
    ok = validate_entry("F1", 1, 2.0, 8.5, "Morning") and validate_entry("F1", 1, 7.0, 8.5, "Evening")
    rows.append(("Both accepted" if ok else "Not accepted", status(ok)))
    ok = rejected("F1", -1, 3.5, 8.5, "Morning") and rejected("F1", 1, 7.1, 8.5, "Morning") \
        and rejected("F1", 1, 3.5, 8.5, "Night")
    rows.append(("All three rejected" if ok else "Not rejected", status(ok)))
    a, b = ten_day_payment(10, 3.5, 8.5, 4, 500), ten_day_payment(10, 3.5, 8.5, 4, 501)
    rows.append((f"500 L: ₹{a} (no bonus); 501 L: ₹{b} (bonus)",
                 status(a == Decimal("320.00") and b == Decimal("325.00"))))
    return rows


def table_timing():
    rows = []
    for n in (300, 6000, 150000):
        t0 = time.perf_counter(); float_total(n)
        t1 = time.perf_counter(); decimal_total(n)
        t2 = time.perf_counter()
        rows.append((f"{t1 - t0:.6f} s", f"{t2 - t1:.6f} s"))
    return rows


def table_seasonal():
    j, a = rate_per_litre("4.2", "8.7", 1), rate_per_litre("4.2", "8.7", 4)
    p500, p501 = ten_day_payment(10, 3.5, 8.5, 4, 500), ten_day_payment(10, 3.5, 8.5, 4, 501)
    return [(f"₹{j}/L", status(j == Decimal("33.40"))),
            (f"₹{a}/L", status(a == Decimal("35.75"))),
            (f"₹{p500}", status(p500 == Decimal("320.00"))),
            (f"₹{p501}", status(p501 == Decimal("325.00")))]


def set_text(cell, text):
    para = cell.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def main():
    src = os.path.join(HERE, "assignment_report.docx")
    out = os.path.join(HERE, "assignment_report_filled.docx")
    doc = Document(src)

    for p in doc.paragraphs:
        t = p.text
        new = None
        if t.startswith("Student name:"):
            new = f"Student name: {NAME}"
        elif t.startswith("Register number:"):
            new = f"Register number: {REG}"
        elif t.startswith("Repository:"):
            new = f"Repository: {REPO}"
        elif t.startswith("Note: This report deliberately"):
            new = ("Note: Results and timings in this report were produced by running fill_report.py "
                   f"(Python {platform.python_version()}, {platform.system()}). Timings vary by computer. "
                   "Screenshots and the GitHub URL are added separately.")
        if new and p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""

    t0, t1, t2 = doc.tables
    for row, (actual, st) in zip(t0.rows[1:], table_tests()):
        set_text(row.cells[2], actual); set_text(row.cells[3], st)
    for row, (f, d) in zip(t1.rows[1:], table_timing()):
        set_text(row.cells[2], f); set_text(row.cells[3], d)
    for row, (actual, st) in zip(t2.rows[1:], table_seasonal()):
        set_text(row.cells[2], actual); set_text(row.cells[3], st)

    doc.save(out)
    print("Saved:", out)


if __name__ == "__main__":
    main()
